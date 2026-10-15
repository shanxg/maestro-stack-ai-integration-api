from dotenv import load_dotenv
load_dotenv()

import os
import sys
import requests
from openpyxl import load_workbook

# ============================================================
# SETTINGS
# ============================================================
GITHUB_TOKEN = os.getenv("GITHUB_TOKEN")
GITHUB_PROJECT_TOKEN = os.getenv("GITHUB_PROJECT_TOKEN")  # classic token, read:project only
GITHUB_REPO  = "shanxg/maestro-stack-ai-integration-api"   # owner/repo
EXCEL_FILE   = "security_plan_maestro_api.xlsx"
SHEET_NAME   = "Security Remediation Backlog"
HEADER_ROW   = 1  # Row that holds the column headers
PROJECT_OWNER  = "shanxg"
PROJECT_NUMBER = 2  # https://github.com/users/shanxg/projects/2


# Column map (1-based indexes, as used by openpyxl)
COL = {
    "ID":             1,
    "PRIORITY":       2,
    "CATEGORY":       3,
    "IMPROVEMENT":    4,
    "LOCATION":       5,
    "FIX":            6,
    "REFERENCE":      7,
    "ALIGNMENT":      8,
    "DECISION":       9,
    "EFFORT":         10,
    "IMPACT":         11,
    "STATUS":         12,
    "EVIDENCE":       13,
    "ISSUE_NUMBER":   14,
    "ISSUE_URL":      15,
}

HEADERS = {
    "Authorization": f"Bearer {GITHUB_TOKEN}",
    "Accept": "application/vnd.github+json",
    "X-GitHub-Api-Version": "2022-11-28",
}

# ============================================================
# HELPERS
# ============================================================
def normalize_label(text):
    """Turns 'Front-end / session' into 'front-end-session'."""
    if not text:
        return None
    import unicodedata
    text = unicodedata.normalize("NFKD", str(text))
    text = text.encode("ascii", "ignore").decode("ascii")
    text = text.lower().replace(" ", "-").replace("/", "-")
    return "-".join(filter(None, text.split("-")))[:50]

def build_title(id_, improvement):
    """E.g. '[1.01] JWT stored in localStorage, readable by any XSS'."""
    title = f"[{id_}] {improvement or 'Untitled'}"
    return title[:250]  # GitHub caps titles at 256 characters

def build_body(row):
    """Builds the issue body in Markdown from the relevant columns."""
    def v(col_name):
        val = row[COL[col_name] - 1].value
        return str(val).strip() if val is not None else "—"

    return f"""## 📌 Plan item
- **ID:** {v("ID")}
- **Priority:** {v("PRIORITY")}
- **Category:** {v("CATEGORY")}

## 🔍 Finding
- **Decision:** {v("DECISION")}
- **Improvement:** {v("IMPROVEMENT")}


**Location:** `{v("LOCATION")}`

## 🛠️ Proposed fix
{v("FIX")}

## 📚 Reference
{v("REFERENCE")}

## 🎯 Role alignment
{v("ALIGNMENT")}

## ⚖️ Estimates
- **Effort:** {v("EFFORT")}
- **Impact:** {v("IMPACT")}

---
_Issue created automatically from the `{EXCEL_FILE}` spreadsheet._
"""

def build_labels(row):
    """Builds the issue labels from the spreadsheet."""
    labels   = ["backlog"]
    priority = row[COL["PRIORITY"] - 1].value
    category = row[COL["CATEGORY"] - 1].value
    effort   = row[COL["EFFORT"] - 1].value
    impact   = row[COL["IMPACT"] - 1].value

    if priority: labels.append(f"priority:{str(priority).lower()}")
    if category: labels.append(normalize_label(category))
    if effort:   labels.append(f"effort:{normalize_label(effort)}")
    if impact:   labels.append(f"impact:{normalize_label(impact)}")
    return labels

# ============================================================
# GITHUB API
# ============================================================
def create_issue(title, body, labels):
    url = f"https://api.github.com/repos/{GITHUB_REPO}/issues"
    payload = {"title": title, "body": body, "labels": labels}
    r = requests.post(url, headers=HEADERS, json=payload, timeout=30)
    if r.status_code == 201:
        data = r.json()
        return data["number"], data["html_url"]
    print(f"   ❌ Error {r.status_code}: {r.text[:200]}")
    return None, None

PROJECT_STATUS_QUERY = """
query($login: String!, $number: Int!, $field: String!) {
    user(login: $login) {
        projectV2(number: $number) {
            items(first: 100) {
                pageInfo { hasNextPage }
                nodes {
                    content{ ... on Issue {number}}
                    fieldValueByName(name: $field) { ... on ProjectV2ItemFieldSingleSelectValue { name } }
                }
            }
        }
    }
}
"""
def list_project_statuses():
    """Returns {issue_number: status} from the GitHub Projects board."""
    payload = {
        "query": PROJECT_STATUS_QUERY,
        "variables": {"login": PROJECT_OWNER, "number": PROJECT_NUMBER, "field": "Status"},
    }
    headers = {"Authorization": f"Bearer {GITHUB_PROJECT_TOKEN}"}
    r = requests.post("https://api.github.com/graphql", headers=headers, json=payload, timeout=30)
    data = r.json() if r.status_code == 200 else {}
    # GraphQL answers 200 even when the query fails, so the "errors" key must be checked too
    if r.status_code != 200 or "errors" in data:
        print(f"❌ Failed to read the project board: {r.status_code} {r.text[:200]}")
        return {}

    items = data["data"]["user"]["projectV2"]["items"]
    if items["pageInfo"]["hasNextPage"]:
        print("   ⚠️  The board has more than 100 items; only the first 100 were read.")

    statuses = {}
    for node in items["nodes"]:
        number = (node["content"] or {}).get("number")         # PRs and drafts have no issue number
        status = (node["fieldValueByName"] or {}).get("name")  # cards without a Status are skipped
        if number and status:
            statuses[number] = status
    return statuses

# ============================================================
# SYNC
# ============================================================
def excel_to_github():
    """Creates issues for rows whose Issue Number is empty."""
    print("\n🚀 [1/2] Excel → GitHub (create new issues)")
    wb = load_workbook(EXCEL_FILE)
    ws = wb[SHEET_NAME]

    created = 0
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        id_          = ws.cell(r, COL["ID"]).value
        improvement  = ws.cell(r, COL["IMPROVEMENT"]).value
        issue_number = ws.cell(r, COL["ISSUE_NUMBER"]).value
        issue_url    = ws.cell(r, COL["ISSUE_URL"]).value

        # Skip empty rows and rows that are already synced
        if not id_ or not improvement:
            continue
        if issue_number and issue_url:
            continue

        row    = [ws.cell(r, c) for c in range(1, ws.max_column + 1)]
        title  = build_title(id_, improvement)
        body   = build_body(row)
        labels = build_labels(row)

        print(f"   → Creating issue [{id_}]: {str(improvement)[:60]}...")
        number, url = create_issue(title, body, labels)

        if number:
            ws.cell(r, COL["ISSUE_NUMBER"], number)
            ws.cell(r, COL["ISSUE_URL"], url)
            created += 1
        else:
            print(f"   ⚠️  Failed to create the issue for ID {id_}")

    if created:
        wb.save(EXCEL_FILE)
        print(f"   ✅ {created} issue(s) created. Spreadsheet updated.")
    else:
        print("   ℹ️  No new issues to create.")
    wb.close()


def github_to_excel():
    """Updates the Status column (L) from the GitHub Projects board."""
    print("\n🔄 [2/2] GitHub Projects → Excel (sync status)")
    statuses = list_project_statuses()
    if not statuses:
        print("   ℹ️  No status found on the project board.")
        return

    wb = load_workbook(EXCEL_FILE)
    ws = wb[SHEET_NAME]
    updated = 0

    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        number = ws.cell(r, COL["ISSUE_NUMBER"]).value
        if not number:
            continue
        try:
            number = int(number)
        except (TypeError, ValueError):
            continue

        # The board owns Status; rows whose issue is not on the board are left as they are
        board_status = statuses.get(number)
        if board_status and board_status != ws.cell(r, COL["STATUS"]).value:
            ws.cell(r, COL["STATUS"], board_status)
            updated += 1

    if updated:
        wb.save(EXCEL_FILE)
        print(f"   ✅ {updated} row(s) with an updated status.")
    else:
        print("   ℹ️  No status to update.")
    wb.close()


# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    for name, value in (("GITHUB_TOKEN", GITHUB_TOKEN), ("GITHUB_PROJECT_TOKEN", GITHUB_PROJECT_TOKEN)):
        if not value:
            print(f"❌ ERROR: {name} is not set.")
            print("   Add it to the local .env file (gitignored). Avoid exporting it in the")
            print("   shell: it stays in the shell history (see CRED-01 in the remediation log).")
            sys.exit(1)

    if not os.path.exists(EXCEL_FILE):
        print(f"❌ ERROR: file '{EXCEL_FILE}' not found in the current folder.")
        sys.exit(1)

    print("=" * 60)
    print(" SYNC — Maestro Stack AI Integration API")
    print("=" * 60)

    excel_to_github()
    github_to_excel()

    print("\n✨ Sync complete.")
